import pytest
from pathlib import Path
from docx import Document
from src.MindReport.docx_renderer import DocxRenderer, datetimefmt

def test_datetimefmt_filter():
    assert datetimefmt("2026-10-17 12:34:56", "%Y/%m/%d") == "2026/10/17"
    assert datetimefmt("invalid-date", "%Y") == "invalid-date"

def test_docx_renderer_basic_substitution(tmp_path):
    # Create empty docx with placeholder
    tpl_file = tmp_path / "test_tpl.docx"
    doc = Document()
    doc.add_paragraph("Hello {{ Variables.User }}!")
    doc.save(str(tpl_file))

    out_file = tmp_path / "out.docx"
    DocxRenderer.render(str(tpl_file), {"Variables": {"User": "Alice"}}, str(out_file))

    out_doc = Document(str(out_file))
    assert "Hello Alice!" in out_doc.paragraphs[0].text
