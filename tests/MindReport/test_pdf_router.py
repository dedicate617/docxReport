import pytest
from pathlib import Path
from unittest.mock import patch
from src.MindReport.pdf_router import PdfRouter

def test_detect_engine_priority_office():
    with patch.object(PdfRouter, "is_office_available", return_value=True), \
         patch.object(PdfRouter, "is_libreoffice_available", return_value=True), \
         patch.object(PdfRouter, "is_chrome_available", return_value=True):
        assert PdfRouter.detect_engine() == "office"

def test_detect_engine_priority_libreoffice():
    with patch.object(PdfRouter, "is_office_available", return_value=False), \
         patch.object(PdfRouter, "is_libreoffice_available", return_value=True), \
         patch.object(PdfRouter, "is_chrome_available", return_value=True):
        assert PdfRouter.detect_engine() == "libreoffice"

def test_detect_engine_priority_chrome():
    with patch.object(PdfRouter, "is_office_available", return_value=False), \
         patch.object(PdfRouter, "is_libreoffice_available", return_value=False), \
         patch.object(PdfRouter, "is_chrome_available", return_value=True):
        assert PdfRouter.detect_engine() == "chrome"

def test_detect_engine_none():
    with patch.object(PdfRouter, "is_office_available", return_value=False), \
         patch.object(PdfRouter, "is_libreoffice_available", return_value=False), \
         patch.object(PdfRouter, "is_chrome_available", return_value=False):
        assert PdfRouter.detect_engine() is None

def test_convert_fallback_to_chrome(tmp_path):
    docx_file = tmp_path / "dummy.docx"
    docx_file.write_text("dummy")
    pdf_file = tmp_path / "dummy.pdf"

    with patch.object(PdfRouter, "_try_office_convert", return_value=False), \
         patch.object(PdfRouter, "_try_libreoffice_convert", return_value=False), \
         patch.object(PdfRouter, "_try_chrome_convert", return_value=True):
        assert PdfRouter.convert_docx_to_pdf(docx_file, pdf_file) is True

def test_current_system_detects_chrome():
    chrome_bin = PdfRouter.find_chrome_or_edge_bin()
    assert chrome_bin is not None
    assert Path(chrome_bin).exists()
    assert PdfRouter.is_chrome_available() is True
    # On this machine without Office/LibreOffice, detect_engine should be "chrome"
    if not PdfRouter.is_office_available() and not PdfRouter.is_libreoffice_available():
        assert PdfRouter.detect_engine() == "chrome"

def test_real_docx_to_pdf_conversion_via_chrome(tmp_path):
    from docx import Document
    test_docx = tmp_path / "sample.docx"
    test_pdf = tmp_path / "sample.pdf"

    # Create a simple docx
    doc = Document()
    doc.add_heading("Chrome PDF Conversion Test", level=1)
    doc.add_paragraph("This is a paragraph generated for automated testing.")
    doc.save(str(test_docx))

    # Convert using PdfRouter with chrome
    success = PdfRouter.convert_docx_to_pdf(test_docx, test_pdf, force_engine="chrome")
    assert success is True
    assert test_pdf.exists()
    assert test_pdf.stat().st_size > 0
    # Check PDF header
    with open(test_pdf, "rb") as f:
        header = f.read(5)
        assert header == b"%PDF-"


