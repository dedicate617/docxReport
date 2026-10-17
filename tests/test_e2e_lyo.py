import pytest
import sqlite3
from pathlib import Path
from unittest.mock import patch
from docx import Document
from src.MindReport.main import run_report_pipeline
from src.MindReport.pdf_router import PdfRouter

def test_full_pipeline_with_sqlite(tmp_path):
    # 1. Create a mock SQLite database matching LYO_Header schema
    db_file = tmp_path / "mock_batch.db"
    conn = sqlite3.connect(str(db_file))
    conn.execute("CREATE TABLE LYO1_LYO_Header (ProjectNo TEXT, BatchNo TEXT, ProductName TEXT, RecipeName TEXT);")
    conn.execute("INSERT INTO LYO1_LYO_Header VALUES ('P001', 'BATCH_001', 'TestProduct', 'RecipeA');")
    conn.commit()
    conn.close()

    # 2. Mock dsdef
    dsdef_file = tmp_path / "test.dsdef"
    dsdef_file.write_text("""<Report><Tables><Table>
<Name>LYO1_LYO_Header</Name><DataSource>LYO1</DataSource>
<Sql>SELECT ProjectNo, BatchNo, ProductName, RecipeName FROM LYO1_LYO_Header WHERE BatchNo = '[%BatchNo%]'</Sql>
<EnableSampling>False</EnableSampling>
</Table></Tables></Report>""")

    # 3. Mock docx template
    docx_file = tmp_path / "test.docx"
    doc = Document()
    doc.add_paragraph("Product: {{ LYO1_LYO_Header[0].ProductName }}")
    doc.save(str(docx_file))

    pdf_file = tmp_path / "out.pdf"

    # 4. Run pipeline (mocking PDF engine in case LibreOffice/Word is not installed in CI)
    def mock_pdf_convert(docx_p, pdf_p, force_engine="auto"):
        Path(pdf_p).write_bytes(b"%PDF-1.4 dummy")
        return True

    base_dir = Path(__file__).resolve().parent.parent
    with patch.object(PdfRouter, 'convert_docx_to_pdf', side_effect=mock_pdf_convert):
        exit_code = run_report_pipeline(
            base_dir=base_dir,
            docxtpl=str(docx_file),
            dsdef=str(dsdef_file),
            dspara=f"DSName=LYO1;DBType=4;DBPath={db_file}",
            vpara="BatchNo=BATCH_001",
            pdfpath=str(pdf_file)
        )

    assert exit_code == 0
    assert pdf_file.exists()

    autogen_docx = docx_file.with_name(docx_file.stem + "_autogen.docx")
    assert autogen_docx.exists()
    rendered_doc = Document(str(autogen_docx))
    assert "Product: TestProduct" in rendered_doc.paragraphs[0].text
